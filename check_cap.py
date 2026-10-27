from docx import Document
from lxml import etree

doc = Document('期末大论文.docx')
for i, p in enumerate(doc.paragraphs):
    if p.text.strip().startswith("图1 "):
        print(f"段{i}: [{p.text[:50]}]")
        pPr = p._element.pPr
        if pPr is not None:
            print("pPr XML:")
            print(etree.tostring(pPr, pretty_print=True).decode()[:800])
        for r in p.runs:
            rPr = r._element.rPr
            if rPr is not None:
                print("rPr XML:")
                print(etree.tostring(rPr, pretty_print=True).decode()[:800])
        
        # 检查图片段落（前一段）
        prev = doc.paragraphs[i-1]
        print(f"\n前一段(图片) 段{i-1}: [{prev.text[:30]}]")
        pPr2 = prev._element.pPr
        if pPr2 is not None:
            print("pPr XML:")
            print(etree.tostring(pPr2, pretty_print=True).decode()[:800])
        break
